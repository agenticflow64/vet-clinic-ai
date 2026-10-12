"""Конвертация прайс-листа клиники из xlsx в data/services.json.

Источники (только чтение, репозиторий Kashtanka не изменяется):
  1. vet_clinic_pricelist.xlsx — полный прайс (296 услуг, без категорий)
  2. kashtanka.xlsx, лист «Прайс-лист» — короткий список услуг с категориями
     (как на сайте), лист «Категории и услуги» — справочник категорий

Категории для полного прайса назначаются правилами по ключевым словам
(см. app/categorize.py) — детерминированно и покрываемо тестами.

Запуск:
    python scripts/convert_pricelist.py \
        --full /path/to/vet_clinic_pricelist.xlsx \
        --curated /path/to/kashtanka.xlsx \
        --out data/services.json
"""

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.categorize import categorize  # noqa: E402


def norm(s: str) -> str:
    return " ".join(str(s).lower().split())


def load_full(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Прайс-лист"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, price = row[0], row[1]
        if not name:
            continue
        name = str(name).strip()
        price_val = int(price) if price is not None else None
        items.append({"name": name, "price": price_val, "source": "full"})
    return items


def load_curated(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Прайс-лист"]
    items = []
    for cat, name, price in ws.iter_rows(min_row=2, values_only=True):
        if not name:
            continue
        items.append(
            {
                "category": str(cat).strip(),
                "name": str(name).strip(),
                "price": int(price) if price is not None else None,
                "source": "curated",
            }
        )
    return items


def convert(full_xlsx: str, curated_xlsx: str) -> list[dict]:
    full = load_full(full_xlsx)
    curated = load_curated(curated_xlsx)
    curated_names = {norm(i["name"]) for i in curated}

    records = []
    # Полный прайс: категория по правилам
    for item in full:
        records.append(
            {
                "category": categorize(item["name"]),
                "name": item["name"],
                "price": item["price"],
                "source": "full",
            }
        )
    # Короткий список с сайта: добавляем то, чего нет в полном прайсе
    # (имена там «человеческие»: «Стерилизация кошки» вместо
    #  «Овариогистерэктомия кошки» — они хорошо матчатся с запросами)
    added = 0
    for item in curated:
        if norm(item["name"]) in curated_names and norm(item["name"]) not in {
            norm(r["name"]) for r in records
        }:
            records.append(item)
            added += 1

    # стабильные id
    for i, r in enumerate(records, 1):
        r["id"] = i
    return records


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--full", required=True, help="путь к vet_clinic_pricelist.xlsx")
    ap.add_argument("--curated", required=True, help="путь к kashtanka.xlsx")
    ap.add_argument("--out", default="data/services.json")
    args = ap.parse_args()

    records = convert(args.full, args.curated)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    cats: dict[str, int] = {}
    for r in records:
        cats[r["category"]] = cats.get(r["category"], 0) + 1
    print(f"OK: {len(records)} услуг -> {out}")
    for c, n in sorted(cats.items(), key=lambda kv: -kv[1]):
        print(f"  {c}: {n}")


if __name__ == "__main__":
    main()
